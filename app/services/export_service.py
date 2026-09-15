from __future__ import annotations

from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
import logging
from pathlib import Path
from shutil import copy2
from tempfile import NamedTemporaryFile
from typing import Any

from openpyxl import load_workbook

from app.domain.enums import BASE_UNITS
from app.repositories.mongo import MongoRepositories
from app.services.excel_reader import FIELD_MAP, INPUT_SHEET
from app.storage.local import LocalFileStorage

MACHINE_COLUMN = "Machine Filled Fields"
DISCREPANCY_COLUMN = "Discrepancy Flag"
logger = logging.getLogger(__name__)


def _excel_number(value: object) -> object:
    if value is None:
        return None
    try:
        decimal = Decimal(str(value))
    except (InvalidOperation, ValueError):
        return value
    return int(decimal) if decimal == decimal.to_integral() else float(decimal)


def _literal(value: object) -> object:
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


class ExportBlockedError(ValueError):
    pass


class ExportService:
    def __init__(self, repositories: MongoRepositories, storage: LocalFileStorage):
        self.repositories = repositories
        self.storage = storage

    def prepare_export(self, job_id: str) -> None:
        job = self.repositories.get_job(job_id)
        if not job:
            raise ValueError("job not found")
        pending = self.repositories.pending_count(job_id)
        if pending:
            raise ExportBlockedError(f"{pending} review decisions remain pending")
        if job.get("status") == "EXPORTING":
            raise ExportBlockedError("workbook export is already running")
        self.repositories.update_job(job_id, {
            "status": "EXPORTING",
            "error": None,
            "progress.stage": "PREPARING_EXPORT",
        })

    def export_in_background(self, job_id: str) -> None:
        try:
            self.export(job_id)
        except Exception:
            logger.exception("Workbook export failed for job %s", job_id)

    def export(self, job_id: str) -> Path:
        job = self.repositories.get_job(job_id)
        if not job:
            raise ValueError("job not found")
        pending = self.repositories.pending_count(job_id)
        if pending:
            raise ExportBlockedError(f"{pending} review decisions remain pending")

        source = self.storage.get_input_path(job_id)
        with NamedTemporaryFile(suffix=".xlsx", delete=False) as handle:
            temporary = Path(handle.name)
        copy2(source, temporary)
        workbook = None
        try:
            self.repositories.update_job(job_id, {"progress.stage": "LOADING_WORKBOOK"})
            workbook = load_workbook(temporary)
            sheet = workbook[INPUT_SHEET]
            original_sheet_names = tuple(workbook.sheetnames)
            headers = {str(cell.value).strip(): cell.column for cell in sheet[1] if cell.value is not None}
            protected_values = {
                FIELD_MAP[key]: tuple(
                    sheet.cell(row=row_number, column=headers[FIELD_MAP[key]]).value
                    for row_number in range(2, sheet.max_row + 1)
                )
                for key in ("legacy_size", "legacy_uom")
            }
            machine_col = headers.get(MACHINE_COLUMN) or sheet.max_column + 1
            discrepancy_col = headers.get(DISCREPANCY_COLUMN) or max(sheet.max_column + 1, machine_col + 1)
            sheet.cell(1, machine_col, MACHINE_COLUMN)
            sheet.cell(1, discrepancy_col, DISCREPANCY_COLUMN)
            field_columns = {
                "standard_size": headers[FIELD_MAP["standard_size"]],
                "standard_uom": headers[FIELD_MAP["standard_uom"]],
                "standard_pack_size": headers[FIELD_MAP["standard_pack_size"]],
            }
            items = self.repositories.all_items(job_id)
            attention: list[dict[str, Any]] = []
            self.repositories.update_job(job_id, {"progress.stage": "APPLYING_DECISIONS"})
            for item in items:
                row_number = item["row_number"]
                review = item["review"]
                changed: list[str] = []
                if review["overall_status"] == "APPROVED":
                    for field, decision in review["field_decisions"].items():
                        proposed = item["field_proposals"].get(field)
                        if decision == "APPROVED" and proposed is not None:
                            value = _excel_number(proposed) if field != "standard_uom" else proposed
                            if field == "standard_uom" and value not in BASE_UNITS:
                                raise ValueError(f"invalid final UOM at row {row_number}: {value}")
                            sheet.cell(row_number, field_columns[field], value)
                            changed.append(field.removeprefix("standard_").upper())
                elif review["overall_status"] == "OVERRIDDEN":
                    for field, value in (review.get("override_values") or {}).items():
                        if value is not None and field in field_columns:
                            value = _excel_number(value) if field != "standard_uom" else value
                            sheet.cell(row_number, field_columns[field], value)
                            changed.append(field.removeprefix("standard_").upper())
                sheet.cell(row_number, machine_col, "|".join(changed) if changed else "NONE")
                details = item["discrepancy"].get("details") or []
                sheet.cell(row_number, discrepancy_col, "MULTIPLE" if len(details) > 1 else (details[0]["pair"].upper() + "_CONFLICT" if details else "NONE"))
                if item["group"] == "C" or item.get("reason_code") == "NO_RULE" or details or review["overall_status"] in {"REJECTED", "OVERRIDDEN"}:
                    attention.append(item)

            self._write_summary(workbook, job)
            self._write_attention(workbook, attention)
            self._validate_workbook(workbook, original_sheet_names, headers, protected_values)
            self.repositories.update_job(job_id, {"progress.stage": "SAVING_WORKBOOK"})
            workbook.save(temporary)
            workbook.close()
            workbook = None
            storage_key = self.storage.save_output(job_id, temporary)
            self.repositories.update_job(job_id, {
                "status": "EXPORTED", "output_storage_key": storage_key,
                "error": None, "progress.stage": "EXPORTED",
            })
            return self.storage.get_output_path(job_id)
        except Exception as exc:
            self.repositories.update_job(job_id, {
                "status": "FAILED",
                "progress.stage": "FAILED",
                "error": f"{type(exc).__name__}: {exc}",
            })
            raise
        finally:
            if workbook is not None:
                workbook.close()
            temporary.unlink(missing_ok=True)

    @staticmethod
    def _write_summary(workbook: Any, job: dict[str, Any]) -> None:
        if "Cleansing_Run_Summary" in workbook.sheetnames:
            del workbook["Cleansing_Run_Summary"]
        sheet = workbook.create_sheet("Cleansing_Run_Summary")
        rows = [
            ("Job ID", job["job_id"]),
            ("Input filename", job["original_file_name"]),
            ("Ruleset version", job["ruleset_version"]),
            ("Ruleset checksum", job["ruleset_checksum"]),
            ("Departments", ", ".join(job["selected_departments"])),
            ("Exported at", datetime.now(timezone.utc).isoformat()),
        ] + [(key, value) for key, value in (job.get("stats") or {}).items()]
        for row in rows:
            sheet.append(row)

    @staticmethod
    def _write_attention(workbook: Any, items: list[dict[str, Any]]) -> None:
        if "Cleansing_Attention" in workbook.sheetnames:
            del workbook["Cleansing_Attention"]
        sheet = workbook.create_sheet("Cleansing_Attention")
        columns = [
            "Item_no", "Department", "Group", "Reason", "Current Unit Size", "Current UOM",
            "Current Pack Size", "Proposed Unit Size", "Proposed UOM", "Proposed Pack Size",
            "Method", "Rule ID", "Factor", "Evidence", "Confidence", "Discrepancy",
            "Decision", "Reviewer Comment",
        ]
        sheet.append(columns)
        for item in items:
            original = item["original"]
            proposal = item["field_proposals"]
            rule = item.get("rule") or {}
            review = item["review"]
            evidence = "; ".join(f"{e['field']}: {e['fragment']}" for e in item.get("evidence") or [])
            sheet.append([_literal(value) for value in (
                item["item_no"], item["department"], item["group"], item.get("reason_code"),
                original["standard_size"], original["standard_uom"], original["standard_pack_size"],
                proposal["standard_size"], proposal["standard_uom"], proposal["standard_pack_size"],
                item["method"], rule.get("rule_id"), rule.get("factor"), evidence,
                item.get("confidence"), str(item["discrepancy"].get("details") or ""),
                review["overall_status"], review.get("comment"),
            )])

    @staticmethod
    def _validate_workbook(
        workbook: Any,
        original_sheet_names: tuple[str, ...],
        headers: dict[str, int],
        protected_values: dict[str, tuple[object, ...]],
    ) -> None:
        if not set(original_sheet_names).issubset(workbook.sheetnames):
            raise ValueError("an original sheet is missing from output")
        sheet = workbook[INPUT_SHEET]
        for header, expected_values in protected_values.items():
            column = headers[header]
            current_values = tuple(
                sheet.cell(row=row_number, column=column).value
                for row_number in range(2, sheet.max_row + 1)
            )
            if current_values != expected_values:
                raise ValueError(f"protected column {header} changed")
