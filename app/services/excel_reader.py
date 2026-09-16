from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterator

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell

from app.domain.product import InputProduct
from app.services.normalization import blank, clean_text, clean_uom

INPUT_SHEET = "UoM_Field_Extract"

FIELD_MAP = {
    "item_no": "Item_no",
    "division": "Division",
    "department": "Department",
    "category": "Category",
    "subcategory": "Subcategory",
    "section": "Section",
    "legacy_size": "item_size_value",
    "legacy_uom": "item_size_unit",
    "standard_size": "Standardize Unit Size",
    "standard_uom": "Standardize UOM",
    "standard_pack_size": "Standardize Pack Size",
    "web_description_eng": "web_description_eng",
    "web_description_chi": "web_description_chi",
    "item_brand_eng": "item_brand_eng",
    "item_brand_local": "item_brand_local_lang",
    "item_desc_eng": "item_desc_eng",
    "item_desc_local": "item_desc_local_lang",
}

PURGE_HEADERS = {
    "item_brand_eng",
    "item_brand_local_lang",
    "item_desc_eng",
    "item_desc_local_lang",
    "business_unit_no",
    "business_unit_name",
    "department_no",
    "department_name",
    "category_no",
    "category_name",
    "sub_category_no",
    "sub_category_name",
    "section_no",
    "section_name",
}

REQUIRED_HEADERS = set(FIELD_MAP.values()) | PURGE_HEADERS


class WorkbookValidationError(ValueError):
    pass


@dataclass(frozen=True)
class WorkbookRow:
    row_number: int
    raw: dict[str, Any]
    product: InputProduct


def _item_number(cell: Cell) -> str:
    if cell.value is None:
        return ""
    if isinstance(cell.value, int):
        fmt = str(cell.number_format or "")
        if fmt and set(fmt) == {"0"}:
            return f"{cell.value:0{len(fmt)}d}"
    if isinstance(cell.value, float) and cell.value.is_integer():
        return str(int(cell.value))
    return str(cell.value).strip()


def _domain_number(value: Any) -> Decimal | str | None:
    if blank(value):
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    return str(value).strip()


class ExcelReader:
    def iter_rows(self, path: Path) -> Iterator[WorkbookRow]:
        try:
            # Read formulas as formulas. Standardized K/L/M inputs must be fixed
            # values; the Group A validator rejects formulas explicitly instead of
            # silently trusting a cached result.
            workbook = load_workbook(path, read_only=True, data_only=False)
        except Exception as exc:
            raise WorkbookValidationError(f"Workbook cannot be opened: {exc}") from exc
        if INPUT_SHEET not in workbook.sheetnames:
            workbook.close()
            raise WorkbookValidationError(f"Missing required sheet: {INPUT_SHEET}")
        sheet = workbook[INPUT_SHEET]
        cells = next(sheet.iter_rows(min_row=1, max_row=1), ())
        headers = [clean_text(cell.value) for cell in cells]
        present = [header for header in headers if header]
        duplicates = sorted({header for header in present if present.count(header) > 1})
        missing = sorted(REQUIRED_HEADERS - set(present))
        if duplicates or missing:
            workbook.close()
            details = []
            if missing:
                details.append(f"missing headers: {', '.join(missing)}")
            if duplicates:
                details.append(f"duplicate headers: {', '.join(duplicates)}")
            raise WorkbookValidationError("; ".join(details))
        header_index = {header: index for index, header in enumerate(headers) if header}
        try:
            for row_number, row_cells in enumerate(sheet.iter_rows(min_row=2), start=2):
                raw = {
                    header: row_cells[index].value if index < len(row_cells) else None
                    for header, index in header_index.items()
                }
                item_cell = row_cells[header_index[FIELD_MAP["item_no"]]]
                product = InputProduct(
                    row_number=row_number,
                    item_no=_item_number(item_cell),
                    division=clean_text(raw.get(FIELD_MAP["division"])),
                    department=clean_text(raw.get(FIELD_MAP["department"])),
                    category=clean_text(raw.get(FIELD_MAP["category"])),
                    subcategory=clean_text(raw.get(FIELD_MAP["subcategory"])),
                    section=clean_text(raw.get(FIELD_MAP["section"])),
                    legacy_size=_domain_number(raw.get(FIELD_MAP["legacy_size"])),
                    legacy_uom=clean_uom(raw.get(FIELD_MAP["legacy_uom"])),
                    raw_standard_size=raw.get(FIELD_MAP["standard_size"]),
                    raw_standard_uom=raw.get(FIELD_MAP["standard_uom"]),
                    raw_standard_pack_size=raw.get(FIELD_MAP["standard_pack_size"]),
                    standard_size=_domain_number(raw.get(FIELD_MAP["standard_size"])),
                    standard_uom=clean_uom(raw.get(FIELD_MAP["standard_uom"])),
                    standard_pack_size=_domain_number(raw.get(FIELD_MAP["standard_pack_size"])),
                    web_description_eng=clean_text(raw.get(FIELD_MAP["web_description_eng"])),
                    web_description_chi=clean_text(raw.get(FIELD_MAP["web_description_chi"])),
                    item_brand_eng=clean_text(raw.get(FIELD_MAP["item_brand_eng"])),
                    item_brand_local=clean_text(raw.get(FIELD_MAP["item_brand_local"])),
                    item_desc_eng=clean_text(raw.get(FIELD_MAP["item_desc_eng"])),
                    item_desc_local=clean_text(raw.get(FIELD_MAP["item_desc_local"])),
                )
                yield WorkbookRow(row_number=row_number, raw=raw, product=product)
        finally:
            workbook.close()

    def validate(self, path: Path) -> None:
        iterator = self.iter_rows(path)
        try:
            next(iterator, None)
        finally:
            iterator.close()
