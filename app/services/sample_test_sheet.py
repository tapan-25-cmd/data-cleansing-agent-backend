"""The "Sample test" sheet added to the downloaded workbook: a short summary, then one line
per sampled row with the website's, the AI judge's and a person's Yes/No and reason."""
from __future__ import annotations

import re
from typing import Any
from xml.sax.saxutils import escape

from app.services.sample_test_service import SOURCES, summarize

SHEET_NAME = "Sample test"
_NAMES = {"web": "Website check (wellcome.com.hk)", "judge": "AI judge", "person": "Person"}
_BAD_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

COLUMNS = (
    ("Group", 8), ("Item No", 10), ("Product", 34), ("Product (local)", 24), ("Category", 18),
    ("Old size", 12), ("Excel size", 14), ("Tool result", 16), ("Tool action", 22),
    ("Website outcome", 16), ("Website shows", 22), ("Website reason", 46), ("Note", 70), ("Website link", 40),
    ("AI judge Yes/No", 10), ("AI judge reason", 60), ("Person Yes/No", 10), ("Person reason", 46),
)
_OUTCOME_WORDS = {"YES": "Yes", "NO": "No", "NOT_COMPARABLE": "Not comparable"}
_KIND_WORDS = {"TOOL_MISSED": "No · tool missed it", "DATA_MISLED": "No · data misled"}


def _col(i: int) -> str:
    s = ""
    i += 1
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


def _cell(ref: str, value: Any, style: int = 0) -> str:
    st = f' s="{style}"' if style else ""
    if value is None or value == "":
        return f'<c r="{ref}"{st}/>'
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return f'<c r="{ref}"{st}><v>{value}</v></c>'
    text = escape(_BAD_XML.sub("", str(value)))[:32000]
    return f'<c r="{ref}" t="inlineStr"{st}><is><t xml:space="preserve">{text}</t></is></c>'


def _values(v: dict[str, Any] | None) -> str:
    if not v or not v.get("standard_size"):
        return "empty"
    pack = v.get("standard_pack_size")
    return f"{v['standard_size']} {v.get('standard_uom') or ''}".strip() + (f" × {pack}" if pack else "")


def lines_of(test: dict[str, Any], bold: int = 1) -> tuple[list[list[Any]], list[int], int]:
    """The sheet as rows of cells, a style flag per row (``bold`` for headings), and the
    header row's number. Shared by the workbook sheet and the standalone download."""
    s = summarize(test)
    rows_in = test.get("rows") or []
    lines: list[list[Any]] = []
    styles: list[int] = []

    def add(cells: list[Any], style: int = 0) -> None:
        lines.append(cells)
        styles.append(style)

    web_first = test.get("mode") == "web_first"
    by_group = {g: sum(1 for r in rows_in if r["group"] == g) for g in ("A", "B", "C")}
    mix = f"A {by_group['A']}, B {by_group['B']}, C {by_group['C']}"
    if web_first:
        add([f"Sample test · {len(rows_in)} products found on wellcome.com.hk ({mix}) · {test.get('progress', {}).get('scanned', '?')} rows looked up · seed {test.get('seed')}"], bold)
        add(["Products were drawn at random and kept only when the retailer's website has a page with the same item code. "
             "The site's size is compared with the value the workbook carries. Each row is also judged by a second AI model, and a person can add a verdict."])
    else:
        add([f"Sample test · {len(rows_in)} random rows ({mix}) · seed {test.get('seed')}"], bold)
        add(["Each row answered Yes (the tool got it right) or No, with a reason, by three independent sources."])
    add(["How to read the website outcome: Yes = our value is the size the shop sells. No · tool missed it = the site differs and the descriptions "
         "held the clue. No · data misled = the site differs but Excel, the old size and the descriptions all agreed, so nothing in the file "
         "could show it. Not comparable = the site counts pieces and the file holds a weight (or the reverse); both can be true, so the row is "
         "left out of the figure. Rows without an answer are never counted as right. Overall weights each group by its share of all rows."])
    add([])
    add(["Source", "Overall %", "Range (95%)", "Group A %", "A yes/answered", "Group B %", "B yes/answered", "Group C %", "C yes/answered"], bold)
    for block in s["groups"]:
        src = block["source"]
        o = s["overall"].get(src)
        row: list[Any] = [_NAMES[src], o["percent"] if o else "—", f"{o['low']}–{o['high']}" if o else "—"]
        for p in block["per_group"]:
            row += [p["percent"] if p["percent"] is not None else "—", f"{p['yes']}/{p['answered']}"]
        add(row)
    no = s.get("web_no") or {}
    add([f"Website: found {s['web_found']} · No where the tool missed it: {no.get('tool_missed', 0)} · No where the data misled: {no.get('data_misled', 0)} · "
         f"not comparable: {sum(p.get('not_comparable', 0) for p in s['groups'][0]['per_group'])} · "
         f"website and AI judge agree on {s['agreement']['agree']} of {s['agreement']['both']} rows both answered"])
    add([])
    header_row = len(lines) + 1
    add([c for c, _ in COLUMNS], bold)
    for r in rows_in:
        web, judge, person = r.get("web") or {}, r.get("judge") or {}, r.get("person") or {}
        tool = r["action"] if r["group"] == "C" else _values(r.get("final"))
        if r["group"] == "C" and r.get("suggestion"):
            tool += f" (suggests {_values(r['suggestion'])})"
        outcome = _KIND_WORDS.get(web.get("kind") or "") or _OUTCOME_WORDS.get(web.get("answer") or "", "—")
        add([r["group"], r["item_no"], r.get("product"), r.get("product_local"), r.get("category"),
             r.get("legacy"), _values(r.get("excel")), tool, r.get("action"),
             outcome, web.get("spec") or web.get("title"), web.get("reason"), web.get("note"), web.get("url"),
             judge.get("answer") or "—", judge.get("reason"), person.get("answer") or "", person.get("reason")])
    return lines, styles, header_row


def standalone_workbook(test: dict[str, Any]) -> bytes:
    """The sample test as its own small workbook, for the download button."""
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    lines, styles, header_row = lines_of(test)
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    for cells, style in zip(lines, styles):
        ws.append(cells)
        if style:
            for c in ws[ws.max_row]:
                c.font = Font(bold=True)
    for i, (_, w) in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=header_row + 1):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(COLUMNS))}{ws.max_row}"
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def sheet_xml(test: dict[str, Any], bold: int = 0) -> bytes:
    lines, styles, header_row = lines_of(test, bold)

    body = []
    for i, (cells, style) in enumerate(zip(lines, styles), 1):
        body.append(f'<row r="{i}">' + "".join(_cell(f"{_col(j)}{i}", v, style) for j, v in enumerate(cells)) + "</row>")
    cols = "".join(f'<col min="{i}" max="{i}" width="{w}" customWidth="1"/>' for i, (_, w) in enumerate(COLUMNS, 1))
    last = f"{_col(len(COLUMNS) - 1)}{len(lines)}"
    return ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            f'<dimension ref="A1:{last}"/>'
            f'<sheetViews><sheetView workbookViewId="0"><pane ySplit="{header_row}" topLeftCell="A{header_row + 1}" activePane="bottomLeft" state="frozen"/></sheetView></sheetViews>'
            f'<cols>{cols}</cols><sheetData>{"".join(body)}</sheetData>'
            f'<autoFilter ref="A{header_row}:{last}"/></worksheet>').encode("utf-8")


def add_sheet(parts: dict[str, bytes], sheet: bytes, *, sheet_name: str = SHEET_NAME,
              file_prefix: str = "sample_test") -> dict[str, bytes]:
    """Add a worksheet to an xlsx given its workbook, relationships and content-types parts.
    Returns the parts to write (the three updated parts and the new sheet)."""
    workbook = parts["xl/workbook.xml"].decode("utf-8")
    rels = parts["xl/_rels/workbook.xml.rels"].decode("utf-8")
    types = parts["[Content_Types].xml"].decode("utf-8")
    n = 1
    while f"xl/worksheets/{file_prefix}{n}.xml" in parts or f"worksheets/{file_prefix}{n}.xml" in rels:
        n += 1
    target = f"worksheets/{file_prefix}{n}.xml"
    ids = [int(x) for x in re.findall(r'Id="rId(\d+)"', rels)]
    rid = f"rId{max(ids, default=0) + 1}"
    sheet_ids = [int(x) for x in re.findall(r'sheetId="(\d+)"', workbook)]
    sheet_id = max(sheet_ids, default=0) + 1
    rel = (f'<Relationship Id="{rid}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" '
           f'Target="{target}"/>')
    rels = rels.replace("</Relationships>", rel + "</Relationships>")
    prefix = re.search(r"<(\w+:)?sheets[ >]", workbook)
    ns = prefix.group(1) or "" if prefix else ""
    r_prefix = re.search(r'xmlns:(\w+)="http://schemas.openxmlformats.org/officeDocument/2006/relationships"', workbook)
    rp = r_prefix.group(1) if r_prefix else "r"
    entry = f'<{ns}sheet name="{sheet_name}" sheetId="{sheet_id}" {rp}:id="{rid}"/>'
    workbook = workbook.replace(f"</{ns}sheets>", entry + f"</{ns}sheets>", 1)
    override = (f'<Override PartName="/xl/{target}" '
                'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>')
    types = types.replace("</Types>", override + "</Types>")
    return {"xl/workbook.xml": workbook.encode("utf-8"), "xl/_rels/workbook.xml.rels": rels.encode("utf-8"),
            "[Content_Types].xml": types.encode("utf-8"), f"xl/{target}": sheet}
