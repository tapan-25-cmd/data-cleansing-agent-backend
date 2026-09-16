from pathlib import Path

import pytest
from openpyxl import Workbook

from app.services.excel_reader import ExcelReader, FIELD_MAP, INPUT_SHEET, PURGE_HEADERS, WorkbookValidationError


def make_workbook(path: Path, missing: str | None = None) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = INPUT_SHEET
    headers = sorted((set(FIELD_MAP.values()) | PURGE_HEADERS) - ({missing} if missing else set()))
    sheet.append(headers)
    values = {header: None for header in headers}
    values.update({"Item_no": 123, "Department": "03_Grocery 2", "item_size_value": 1, "item_size_unit": "KG"})
    sheet.append([values[header] for header in headers])
    item_cell = sheet.cell(2, headers.index("Item_no") + 1)
    item_cell.number_format = "000000"
    workbook.save(path)


def test_reader_maps_by_header_and_preserves_leading_zero_item(tmp_path: Path):
    path = tmp_path / "input.xlsx"
    make_workbook(path)
    row = next(ExcelReader().iter_rows(path))
    assert row.product.item_no == "000123"
    assert row.product.legacy_uom == "KG"


def test_reader_preserves_raw_standardized_values_for_validation(tmp_path: Path):
    path = tmp_path / "input.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = INPUT_SHEET
    headers = sorted(set(FIELD_MAP.values()) | PURGE_HEADERS)
    sheet.append(headers)
    values = {header: None for header in headers}
    values.update({
        "Item_no": "000123",
        "Department": "03_Grocery 2",
        "Standardize Unit Size": "500",
        "Standardize UOM": " ml ",
        "Standardize Pack Size": "1",
    })
    sheet.append([values[header] for header in headers])
    workbook.save(path)

    row = next(ExcelReader().iter_rows(path))
    assert row.product.raw_standard_size == "500"
    assert row.product.raw_standard_uom == " ml "
    assert row.product.raw_standard_pack_size == "1"
    assert row.product.standard_uom == "ML"


def test_missing_header_fails_before_processing(tmp_path: Path):
    path = tmp_path / "input.xlsx"
    make_workbook(path, missing="Standardize UOM")
    with pytest.raises(WorkbookValidationError, match="Standardize UOM"):
        list(ExcelReader().iter_rows(path))
