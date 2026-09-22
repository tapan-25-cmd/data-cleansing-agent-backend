import re
from io import BytesIO
from zipfile import ZipFile
from uuid import uuid4

from openpyxl import Workbook, load_workbook

from app.services.excel_reader import FIELD_MAP, INPUT_SHEET
from app.services.export_service import AUDIT_COLUMNS, STATUS_FILLS, ExportService
from app.storage.local import LocalFileStorage


class RepositoriesStub:
    def __init__(self) -> None:
        self.job = {"job_id": "test", "status": "READY_FOR_REVIEW"}
        self.updates: list[dict] = []
        self.items = [{
            "row_number": 2,
            "field_proposals": {"standard_size": "1000", "standard_uom": "GM", "standard_pack_size": "2"},
            "review": {"overall_status": "PENDING", "override_values": None},
        }]

    def get_job(self, _job_id: str) -> dict:
        return self.job

    def update_job(self, _job_id: str, values: dict) -> None:
        self.updates.append(values)
        self.job.update({key: value for key, value in values.items() if "." not in key})

    def all_items(self, _job_id: str) -> list[dict]:
        return self.items


def test_export_applies_pending_proposals_without_rebuilding_workbook(tmp_path):
    job_id = str(uuid4())
    storage = LocalFileStorage(tmp_path, 10_000_000)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = INPUT_SHEET
    headers = list(FIELD_MAP.values())
    sheet.append(headers)
    row = [None] * len(headers)
    row[headers.index(FIELD_MAP["item_no"])] = "SKU-1"
    row[headers.index(FIELD_MAP["legacy_size"])] = 1
    row[headers.index(FIELD_MAP["legacy_uom"])] = "KG"
    row[headers.index(FIELD_MAP["standard_size"])] = 1
    row[headers.index(FIELD_MAP["standard_uom"])] = "KG"
    row[headers.index(FIELD_MAP["standard_pack_size"])] = 1
    sheet.append(row)
    validated_row = [None] * len(headers)
    validated_row[headers.index(FIELD_MAP["item_no"])] = "SKU-A"
    validated_row[headers.index(FIELD_MAP["standard_size"])] = 500.5
    validated_row[headers.index(FIELD_MAP["standard_uom"])] = "ML"
    validated_row[headers.index(FIELD_MAP["standard_pack_size"])] = 1
    sheet.append(validated_row)
    source = BytesIO()
    workbook.save(source)
    source.seek(0)
    storage.save_input(job_id, source)
    repositories = RepositoriesStub()
    service = ExportService(repositories, storage)

    service.prepare_export(job_id)
    output = service.export(job_id)

    result = load_workbook(output, read_only=True, data_only=True)
    output_sheet = result[INPUT_SHEET]
    output_headers = {cell.value: cell.column for cell in output_sheet[1]}
    assert output_sheet.cell(2, output_headers[FIELD_MAP["legacy_size"]]).value == 1
    assert output_sheet.cell(2, output_headers[FIELD_MAP["legacy_uom"]]).value == "KG"
    assert output_sheet.cell(2, output_headers[FIELD_MAP["standard_size"]]).value == 1000
    assert output_sheet.cell(2, output_headers[FIELD_MAP["standard_uom"]]).value == "GM"
    assert output_sheet.cell(2, output_headers[FIELD_MAP["standard_pack_size"]]).value == 2
    # Export patches only proposal rows. A validated Group A row remains byte-for-byte
    # equivalent at the cell-value level, including its existing decimal size.
    assert output_sheet.cell(3, output_headers[FIELD_MAP["standard_size"]]).value == 500.5
    assert output_sheet.cell(3, output_headers[FIELD_MAP["standard_uom"]]).value == "ML"
    assert output_sheet.cell(3, output_headers[FIELD_MAP["standard_pack_size"]]).value == 1
    assert result.sheetnames == [INPUT_SHEET]
    assert set(AUDIT_COLUMNS).issubset(output_headers)
    assert output_sheet.cell(2, output_headers["Cleansing Status"]).value == "Already correct"
    assert output_sheet.cell(2, output_headers["Group"]).value == ""
    assert output_sheet.cell(2, output_headers["Comment"]).value == (
        "Checked. The existing values passed every check."
    )
    # The three columns that say what happened sit together, in reading order.
    assert list(AUDIT_COLUMNS[:3]) == ["Group", "Cleansing Status", "Comment"]
    result.close()
    assert repositories.job["status"] == "EXPORTED"


def test_export_does_not_apply_unapproved_review_required_proposal(tmp_path):
    job_id = str(uuid4())
    storage = LocalFileStorage(tmp_path, 10_000_000)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = INPUT_SHEET
    headers = list(FIELD_MAP.values())
    sheet.append(headers)
    row = [None] * len(headers)
    row[headers.index(FIELD_MAP["item_no"])] = "044305"
    row[headers.index(FIELD_MAP["standard_size"])] = 375
    row[headers.index(FIELD_MAP["standard_uom"])] = "GM"
    row[headers.index(FIELD_MAP["standard_pack_size"])] = 1
    sheet.append(row)
    source = BytesIO()
    workbook.save(source)
    source.seek(0)
    storage.save_input(job_id, source)

    repositories = RepositoriesStub()
    repositories.items = [{
        "row_number": 2,
        "item_no": "044305",
        "application_policy": "REVIEW_REQUIRED",
        "field_proposals": {
            "standard_size": "349",
            "standard_uom": "GM",
            "standard_pack_size": None,
        },
        "findings": [{
            "code": "SIGNIFICANT_LEGACY_SIZE_MISMATCH",
            "category": "SOURCE_DISCREPANCY",
            "human_reason": "Existing 375 GM differs from converted 349 GM",
        }],
        "changes": [{
            "field": "standard_size",
            "original": "375",
            "proposed": "349",
            "final": "375",
        }, {
            "field": "standard_uom",
            "original": "GM",
            "proposed": "GM",
            "final": "GM",
        }, {
            "field": "standard_pack_size",
            "original": "1",
            "proposed": None,
            "final": "1",
        }],
        "method": "RULE",
        "result_ledger_version": "result-ledger-v1",
        "review": {
            "overall_status": "PENDING",
            "override_values": None,
            "comment": None,
        },
    }]
    service = ExportService(repositories, storage)

    service.prepare_export(job_id)
    output = service.export(job_id)

    result = load_workbook(output, read_only=True, data_only=True)
    output_sheet = result[INPUT_SHEET]
    output_headers = {cell.value: cell.column for cell in output_sheet[1]}
    assert output_sheet.cell(2, output_headers[FIELD_MAP["standard_size"]]).value == 375
    assert output_sheet.cell(2, output_headers["Proposed K"]).value == "349"
    assert output_sheet.cell(2, output_headers["Final K"]).value == "375"
    assert output_sheet.cell(2, output_headers["Review Status"]).value == "PENDING"
    result.close()


EXCEL_ROOT = (
    '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
    ' xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
    ' xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"'
    ' mc:Ignorable="x14ac xr xr2 xr3"'
    ' xmlns:x14ac="http://schemas.microsoft.com/office/spreadsheetml/2009/9/ac"'
    ' xmlns:xr="http://schemas.microsoft.com/office/spreadsheetml/2014/revision"'
    ' xmlns:xr2="http://schemas.microsoft.com/office/spreadsheetml/2015/revision2"'
    ' xmlns:xr3="http://schemas.microsoft.com/office/spreadsheetml/2016/revision3"'
    ' xr:uid="{EF2B22FA-DBBE-4D6A-BEE0-2ADA851D6098}">'
)


def _as_excel_writes_it(source: BytesIO) -> BytesIO:
    """Give an openpyxl workbook the namespace header a real Excel file carries."""
    source.seek(0)
    original = ZipFile(source)
    rebuilt = BytesIO()
    with ZipFile(rebuilt, "w") as target:
        for entry in original.infolist():
            content = original.read(entry.filename)
            if entry.filename.startswith("xl/worksheets/sheet"):
                content = re.sub(rb"<worksheet\b[^>]*>", EXCEL_ROOT.encode(), content, count=1)
                assert b"mc:Ignorable" in content
            target.writestr(entry, content)
    rebuilt.seek(0)
    return rebuilt


def test_exported_workbook_keeps_the_namespaces_excel_requires(tmp_path):
    """Excel rejects a sheet whose mc:Ignorable names an undeclared prefix.

    openpyxl reads such a file happily, so this asserts the XML itself rather than
    relying on a successful load.
    """
    job_id = str(uuid4())
    storage = LocalFileStorage(tmp_path, 10_000_000)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = INPUT_SHEET
    headers = list(FIELD_MAP.values())
    sheet.append(headers)
    row = [None] * len(headers)
    row[headers.index(FIELD_MAP["item_no"])] = "SKU-1"
    row[headers.index(FIELD_MAP["standard_size"])] = 1
    row[headers.index(FIELD_MAP["standard_uom"])] = "KG"
    row[headers.index(FIELD_MAP["standard_pack_size"])] = 1
    sheet.append(row)
    source = BytesIO()
    workbook.save(source)
    storage.save_input(job_id, _as_excel_writes_it(source))

    service = ExportService(RepositoriesStub(), storage)
    output = service.export(job_id)

    with ZipFile(output) as archive:
        name = next(n for n in archive.namelist() if n.startswith("xl/worksheets/sheet"))
        worksheet = archive.read(name)
    root = re.search(rb"<worksheet\b[^>]*>", worksheet).group(0).decode()

    declared = set(re.findall(r'xmlns:([\w.-]+)=', root))
    ignorable = set(re.search(r'Ignorable="([^"]*)"', root).group(1).split())
    assert ignorable and ignorable <= declared, (
        f"Excel would call this workbook damaged: mc:Ignorable names "
        f"{sorted(ignorable - declared)}, which nothing declares"
    )
    # The workbook's own prefixes survive; nothing is renamed to ns0/ns1/...
    assert {"r", "mc", "x14ac", "xr", "xr2", "xr3"} <= declared
    assert not re.search(r'xmlns:ns\d', root)
    assert "xr:uid" in root


def test_status_cell_is_coloured_and_the_style_indexes_resolve(tmp_path):
    """A style index that points past the end of styles.xml also makes Excel
    call the workbook damaged, so check the indexes rather than just the colour."""
    job_id = str(uuid4())
    storage = LocalFileStorage(tmp_path, 10_000_000)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = INPUT_SHEET
    headers = list(FIELD_MAP.values())
    sheet.append(headers)
    row = [None] * len(headers)
    row[headers.index(FIELD_MAP["item_no"])] = "SKU-1"
    row[headers.index(FIELD_MAP["standard_size"])] = 1
    row[headers.index(FIELD_MAP["standard_uom"])] = "KG"
    row[headers.index(FIELD_MAP["standard_pack_size"])] = 1
    sheet.append(row)
    source = BytesIO()
    workbook.save(source)
    source.seek(0)
    storage.save_input(job_id, source)

    output = ExportService(RepositoriesStub(), storage).export(job_id)

    with ZipFile(output) as archive:
        styles = archive.read("xl/styles.xml").decode()
        name = next(n for n in archive.namelist() if n.startswith("xl/worksheets/sheet"))
        worksheet = archive.read(name).decode()

    fills = re.findall(r"<fill>.*?</fill>", re.search(r"<fills.*?</fills>", styles, re.S).group(0), re.S)
    records = re.findall(r"<xf [^>]*/>", re.search(r"<cellXfs.*?</cellXfs>", styles, re.S).group(0))
    assert int(re.search(r'<fills count="(\d+)"', styles).group(1)) == len(fills)
    assert int(re.search(r'<cellXfs count="(\d+)"', styles).group(1)) == len(records)

    status_cell = re.search(r'<c r="[A-Z]+2"[^>]*s="(\d+)"[^>]*>', worksheet)
    assert status_cell, "the status cell carries no style"
    style_index = int(status_cell.group(1))
    assert style_index < len(records), "style index points past the end of styles.xml"
    fill_index = int(re.search(r'fillId="(\d+)"', records[style_index]).group(1))
    assert fill_index < len(fills), "fill index points past the end of styles.xml"
    colour = re.search(r'rgb="([0-9A-F]{8})"', fills[fill_index]).group(1)
    # This row is already correct, so it takes the green fill.
    assert colour == STATUS_FILLS["NO_CHANGE"]
    # Every status has its own colour; none collide.
    assert len(set(STATUS_FILLS.values())) == len(STATUS_FILLS)


def test_the_comment_says_what_happened_to_each_kind_of_row():
    """Every row explains itself: a purged row must not claim it was checked."""
    from app.services.export_service import _comment

    corrected = _comment({
        "group": "B", "application_policy": "AUTO_APPLY", "method": "RULE",
        "original": {"legacy_size": "16", "legacy_uom": "OZ"},
        "field_proposals": {"standard_size": "454", "standard_uom": "GM"},
        "changes": [
            {"field": "standard_size", "original": None, "proposed": "454"},
            {"field": "standard_uom", "original": "OZ", "proposed": "GM"},
        ],
        "findings": [],
    })
    assert corrected == (
        "Corrected from the legacy data (16 OZ): size blank → 454, unit OZ → GM."
    )

    purged = _comment({"group": "SKIPPED_PURGED", "application_policy": "NO_CHANGE", "findings": []})
    assert purged.startswith("This product record is empty")
    assert "passed every check" not in purged

    no_rule = _comment({"group": "B", "reason_code": "NO_RULE", "application_policy": "NO_CHANGE",
                        "original": {"legacy_size": "1", "legacy_uom": "ST"}, "findings": [],
                        "field_proposals": {"standard_size": None, "standard_uom": None}})
    assert no_rule == ("The legacy unit ST has no agreed conversion, so the size could not be "
                       "filled in. It was left blank rather than guessed.")

    untouched = _comment({"group": "A", "application_policy": "NO_CHANGE", "findings": []})
    assert untouched == "Checked. The existing values passed every check."


def test_review_comments_carry_the_numbers_a_person_needs():
    from app.services.export_service import _comment

    row = {
        "group": "A", "application_policy": "REVIEW_REQUIRED",
        "original": {"legacy_size": "12.3", "legacy_uom": "OZ", "standard_size": 375, "standard_uom": "GM"},
        "findings": [{
            "code": "SIGNIFICANT_LEGACY_SIZE_MISMATCH", "field": "standard_size",
            "human_reason": "Existing standardized size differs from the rounded legacy conversion",
            "evidence": [{"role": "CURRENT", "value": "375"}, {"role": "EXPECTED", "value": "349"}],
        }],
    }
    assert _comment(row) == (
        "Excel says 375 GM, but the legacy data (12.3 OZ) works out to 349 GM. "
        "Check the pack and confirm which is right."
    )
    # Same unit on both sides: nothing was converted, so do not say it was.
    same_unit = {**row, "original": {**row["original"], "legacy_size": "450", "legacy_uom": "GM"}}
    same_unit["findings"] = [{**row["findings"][0], "evidence": [
        {"role": "CURRENT", "value": "18"}, {"role": "EXPECTED", "value": "450"},
    ]}]
    assert _comment(same_unit).startswith(
        "Excel says 18 GM, but the legacy data says 450 GM."
    )

    mixed_units = {
        "group": "A", "application_policy": "REVIEW_REQUIRED",
        "original": {"legacy_size": "100", "legacy_uom": "ML", "standard_uom": "GM"},
        "findings": [{
            "code": "LEGACY_UOM_MISMATCH", "field": "standard_uom", "human_reason": "jargon",
            "evidence": [{"role": "CURRENT", "value": "GM"}, {"role": "EXPECTED", "value": "ML"}],
        }],
    }
    assert _comment(mixed_units) == (
        "Excel measures this in GM (a weight), but the legacy data says 100 ML, which is "
        "a volume. These measure different things, so confirm which one applies."
    )
    comment = _comment(mixed_units)
    # No internal vocabulary reaches the spreadsheet.
    for jargon in ("standardized", "K/L/M", "UOM", "Group A", "deterministic"):
        assert jargon not in comment


def test_a_job_exported_by_an_older_version_is_rebuilt_not_served_stale(tmp_path):
    """A fix to the exporter has to reach jobs that were already downloaded once."""
    from app.services.export_service import EXPORT_VERSION

    job_id = str(uuid4())
    storage = LocalFileStorage(tmp_path, 10_000_000)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = INPUT_SHEET
    headers = list(FIELD_MAP.values())
    sheet.append(headers)
    row = [None] * len(headers)
    row[headers.index(FIELD_MAP["item_no"])] = "SKU-1"
    sheet.append(row)
    source = BytesIO()
    workbook.save(source)
    source.seek(0)
    storage.save_input(job_id, source)

    repositories = RepositoriesStub()
    service = ExportService(repositories, storage)
    service.export(job_id)
    assert repositories.job["export_version"] == EXPORT_VERSION

    # Same version: the file on disk is reused.
    stamped = storage.get_output_path(job_id).stat().st_mtime_ns
    service.export(job_id)
    assert storage.get_output_path(job_id).stat().st_mtime_ns == stamped

    # Built by an older exporter: rebuild rather than hand back the stale file.
    repositories.job["export_version"] = "export-v1"
    service.export(job_id)
    assert storage.get_output_path(job_id).stat().st_mtime_ns != stamped
    assert repositories.job["export_version"] == EXPORT_VERSION


def test_the_autofilter_is_widened_so_sorting_in_excel_keeps_the_audit_columns_with_their_rows(tmp_path):
    """Excel sorts only the AutoFilter range. A filter left over the original columns
    means one sort in Excel scrambles every status against its row."""
    job_id = str(uuid4())
    storage = LocalFileStorage(tmp_path, 10_000_000)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = INPUT_SHEET
    headers = list(FIELD_MAP.values())
    sheet.append(headers)
    row = [None] * len(headers)
    row[headers.index(FIELD_MAP["item_no"])] = "SKU-1"
    sheet.append(row)
    sheet.auto_filter.ref = f"A1:{sheet.cell(1, len(headers)).column_letter}2"
    source = BytesIO()
    workbook.save(source)
    source.seek(0)
    storage.save_input(job_id, source)

    output = ExportService(RepositoriesStub(), storage).export(job_id)

    with ZipFile(output) as archive:
        name = next(n for n in archive.namelist() if n.startswith("xl/worksheets/sheet"))
        worksheet = archive.read(name).decode()
    dimension = re.search(r'<dimension ref="A1:([A-Z]+)\d+"', worksheet).group(1)
    auto_filter = re.search(r'<autoFilter ref="A1:([A-Z]+)\d+"', worksheet).group(1)
    assert auto_filter == dimension, "AutoFilter must cover the appended audit columns"
    last_audit = re.findall(r'<c r="([A-Z]+)1"', worksheet)[-1]
    assert auto_filter == last_audit
