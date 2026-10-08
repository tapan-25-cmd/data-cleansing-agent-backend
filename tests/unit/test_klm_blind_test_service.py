"""K L M blind test: a blind checker works out the right values; plain rules compare them with ours."""
import json
import re

from app.agents.klm_checker import CheckRequest, CheckResult, Evidence, MockKlmChecker, validate_check
from app.agents.judge import Values
from app.services.klm_blind_sheet import COLUMNS, REVIEWER_HEADER, lines_of, sheet_name, sheet_xml
from app.services.klm_blind_test_service import KlmBlindTestService, check_request, compare, draw, draw_more, same, sample_numbers, summarize

import pytest


def item(n, route, policy, desc="SUGAR 500G", excel=(500, "GM", 1), proposals=None):
    return {"row_number": n, "item_no": f"{n:06d}", "route": route, "application_policy": policy,
            "context": {"item_desc_eng": desc, "category": "Sugar"},
            "original": {"legacy_size": "500", "legacy_uom": "GM", "standard_size": excel[0], "standard_uom": excel[1], "standard_pack_size": excel[2]},
            "field_proposals": proposals or {"standard_size": None, "standard_uom": None, "standard_pack_size": None},
            "findings": [], "changes": [], "review": {"overall_status": "NOT_REQUIRED"}}


def v(size, uom, pack=1):
    return {"standard_size": size, "standard_uom": uom, "standard_pack_size": pack}


def test_the_checker_is_shown_the_raw_data_and_nothing_the_tool_produced():
    changed = item(7, "B", "AUTO_APPLY", excel=(1, "KG", 1), proposals={"standard_size": "1000", "standard_uom": "GM", "standard_pack_size": None})
    changed.update(reason_code="RULE_CONVERSION", findings=[{"code": "CONVERTED"}], reasoner={"outcome": "CONFIRMED"})
    sent = json.loads(check_request(changed).model_dump_json())
    assert set(sent) <= {"task", "descriptions", "category", "subcategory", "category_kind", "legacy", "excel", "repair_attempt", "repair_validation_error"}
    assert sent["excel"] == {"size": "1", "uom": "KG", "pack_size": "1"}, "Excel as uploaded, not what the tool wrote"
    assert "1000" not in json.dumps(sent) and "CONFIRMED" not in json.dumps(sent) and "AUTO_APPLY" not in json.dumps(sent)


def test_the_same_amount_is_the_same_however_it_is_split():
    assert same(v(70, "GM", 3), v(210, "GM", 1)) and same(v(4, "EA", 1), v(1, "EA", 4)) and same(v(1, "KG"), v(1000, "GM"))
    assert not same(v(500, "GM"), v(500, "ML")) and not same(v(500, "GM"), v(450, "GM")) and not same(None, v(1, "EA"))


def test_what_right_means_in_each_group():
    kept = {"excel": v(500, "GM"), "final": v(500, "GM"), "suggestion": None}
    assert compare("A", kept, {"decision": "AS_IS", "values": v(500, "GM")})["verdict"] == "RIGHT"
    missed = compare("A", kept, {"decision": "VALUES", "values": v(450, "GM")})
    assert (missed["verdict"], missed["error_label"]) == ("WRONG", "Missed error") and "450 GM" in missed["why"]
    assert compare("A", kept, {"decision": "NEEDS_PERSON", "values": None})["verdict"] == "CANT_TELL"

    wrote = {"excel": v(1, "KG"), "final": v(1000, "GM"), "suggestion": None}
    assert compare("B", wrote, {"decision": "VALUES", "values": v(1000, "GM")})["verdict"] == "RIGHT"
    assert compare("B", wrote, {"decision": "VALUES", "values": v(100, "GM")})["error_label"] == "Wrong correction"

    raised = {"excel": v(450, "GM", 6), "final": v(450, "GM", 6), "suggestion": v(90, "GM", 6)}
    assert compare("C", raised, {"decision": "NEEDS_PERSON", "values": None})["verdict"] == "RIGHT"
    assert compare("C", raised, {"decision": "VALUES", "values": v(90, "GM", 6)})["verdict"] == "RIGHT"
    assert compare("C", raised, {"decision": "VALUES", "values": v(100, "GM", 6)})["error_label"] == "Wrong suggestion"
    assert compare("C", raised, {"decision": "AS_IS", "values": v(450, "GM", 6)})["error_label"] == "False alarm"
    # A suggestion that only re-splits the same total (90 GM × 30 for 450 GM × 6) is the same amount.
    resplit = {**raised, "suggestion": v(90, "GM", 30)}
    assert compare("C", resplit, {"decision": "AS_IS", "values": v(450, "GM", 6)})["verdict"] == "RIGHT"
    blank = {"excel": v(None, None, None), "final": v(None, None, None), "suggestion": None}
    assert compare("C", blank, {"decision": "VALUES", "values": v(940, "ML")})["verdict"] == "RIGHT"
    # A row the checker could not answer is never counted as right.
    assert compare("A", kept, {"decision": None, "values": None, "error": "TimeoutError"})["verdict"] == "CANT_TELL"


def test_an_answer_must_carry_its_values_and_quote_words_that_are_there():
    request = CheckRequest(descriptions={"item_desc_eng": "JUICE 500ML", "item_brand_eng": "BRAND 7"}, excel=Values(size="500", uom="ML", pack_size="1"))
    validate_check(request, CheckResult(decision="AS_IS", size="500", uom="ML", pack_size="1", reason="", evidence=[Evidence(field="item_desc_eng", fragment="500ML")]))
    with pytest.raises(ValueError):
        validate_check(request, CheckResult(decision="VALUES", reason=""))
    with pytest.raises(ValueError):
        validate_check(request, CheckResult(decision="NEEDS_PERSON", size="500", uom="ML", reason=""))
    with pytest.raises(ValueError):
        validate_check(request, CheckResult(decision="AS_IS", size="500", uom="ML", reason="", evidence=[Evidence(field="item_desc_eng", fragment="750ML")]))
    with pytest.raises(ValueError):
        validate_check(request, CheckResult(decision="AS_IS", size="500", uom="ML", reason="", evidence=[Evidence(field="item_brand_eng", fragment="7")]))


def test_a_run_draws_the_same_rows_for_a_seed_and_gives_a_figure_per_group():
    items = {"A": [item(n, "A", "NO_CHANGE") for n in range(1, 41)] + [item(41, "A", "NO_CHANGE", desc="CHANGE ME 500G")],
             "B": [item(n, "B", "AUTO_APPLY", excel=(500, "GM", None), proposals={"standard_size": None, "standard_uom": None, "standard_pack_size": "1"}) for n in range(50, 70)],
             "C": [item(n, "A", "REVIEW_REQUIRED", desc="PERSON PLEASE 500G") for n in range(80, 95)]}
    sampled = draw(items, 10, seed=3)
    assert [x["row_number"] for x in draw(items, 10, seed=3)["A"]] == [x["row_number"] for x in sampled["A"]]
    assert {g: len(rows) for g, rows in sampled.items()} == {"A": 10, "B": 10, "C": 10}
    full = {"A": items["A"], "B": items["B"], "C": items["C"]}
    rows = KlmBlindTestService(MockKlmChecker(), 3).run(full)
    test = {"job_id": "j", "status": "READY", "size": 41, "seed": 3, "populations": {"A": 900, "B": 60, "C": 40}, "rows": rows, "checker_version": "v"}
    by = {g["group"]: g for g in summarize(test)["groups"]}
    assert (by["A"]["right"], by["A"]["wrong"], by["A"]["errors"]) == (40, 1, {"MISSED_ERROR": 1})
    # We filled the pack as 1; the checker reaches 500 GM × 1 on its own.
    assert by["B"]["right"] == 20 and by["C"]["right"] == 15
    assert summarize(test)["overall"]["checked"] == 76 and summarize(test)["weighted_percent"] is not None

    # The sheet ends with the reviewer's Yes/No, as a drop-down over the data rows.
    rows[0]["person"] = {"answer": "NO", "comment": "pack is 2"}
    lines, _, header_row = lines_of(test)
    assert lines[header_row - 1][-1] == REVIEWER_HEADER == COLUMNS[-1][0]
    assert lines[header_row][-1] == "No" and lines[header_row][-2] == "pack is 2" and lines[header_row + 1][-1] == ""
    xml = sheet_xml(test).decode()
    assert re.search(r'<dataValidation type="list"[^>]*sqref="Y%d:Y%d"' % (header_row + 1, len(lines)), xml) and '"Yes,No"' in xml
    assert summarize(test)["person"]["checked"] == 1
    # Before a run, the sheet says so.
    assert lines_of(None)[0][0] == ["K L M blind test · not run"]


def test_a_finished_test_is_a_sheet_in_the_exported_workbook_and_a_new_answer_rebuilds_it(tmp_path):
    from datetime import datetime, timedelta, timezone
    from io import BytesIO
    from uuid import uuid4

    from openpyxl import Workbook, load_workbook

    from app.services.excel_reader import FIELD_MAP, INPUT_SHEET
    from app.services.export_service import ExportService
    from app.services.klm_blind_sheet import SHEET_NAME
    from app.storage.local import LocalFileStorage

    rows = KlmBlindTestService(MockKlmChecker(), 3).run({"A": [item(2, "A", "NO_CHANGE"), item(3, "A", "NO_CHANGE", desc="CHANGE ME 500G")], "B": [], "C": []})
    rows[1]["person"] = {"answer": "YES", "comment": "the file is right"}

    class Repositories:
        def __init__(self) -> None:
            self.job = {"job_id": "test", "status": "READY_FOR_REVIEW"}
            self.test = {"job_id": "test", "status": "READY", "size": 2, "seed": 9, "populations": {"A": 2, "B": 0, "C": 0},
                         "rows": rows, "checker_version": "v", "saved_at": datetime.now(timezone.utc)}

        def get_job(self, _job_id): return self.job
        def update_job(self, _job_id, values): self.job.update({k: v for k, v in values.items() if "." not in k})
        def all_items(self, _job_id): return []
        def klm_blind_test(self, _job_id): return self.test

    job_id = str(uuid4())
    storage = LocalFileStorage(tmp_path, 10_000_000)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = INPUT_SHEET
    sheet.append(list(FIELD_MAP.values()))
    sheet.append(["SKU-1"])
    source = BytesIO()
    workbook.save(source)
    source.seek(0)
    storage.save_input(job_id, source)

    repositories = Repositories()
    service = ExportService(repositories, storage)
    result = load_workbook(service.export(job_id))
    assert sheet_name(1) in result.sheetnames and SHEET_NAME not in result.sheetnames
    ws = result[sheet_name(1)]
    assert ws["A1"].value.startswith("K L M blind test · 2 rows")
    header = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Group" and ws.cell(r, 2).value == "Excel row")
    last = ws.max_column
    assert ws.cell(header, last).value == REVIEWER_HEADER
    assert [ws.cell(header + 1, c).value for c in (1, 19, last)] == ["A", "Right", None]
    assert [ws.cell(header + 2, c).value for c in (19, 20, last - 1, last)] == ["Wrong", "Missed error", "the file is right", "Yes"]
    validation = ws.data_validations.dataValidation[0]
    assert validation.type == "list" and validation.formula1 == '"Yes,No"' and f"Y{header + 1}" in str(validation.sqref)

    # A reviewer's answer saved after the export makes the file stale; otherwise it is reused.
    stamped = storage.get_output_path(job_id).stat().st_mtime_ns
    service.export(job_id)
    assert storage.get_output_path(job_id).stat().st_mtime_ns == stamped
    repositories.test["saved_at"] = repositories.job["exported_at"] + timedelta(seconds=5)
    service.export(job_id)
    assert storage.get_output_path(job_id).stat().st_mtime_ns != stamped


def test_several_samples_share_no_row_and_each_gets_its_own_sheet():
    items = {"A": [item(n, "A", "NO_CHANGE") for n in range(1, 61)],
             "B": [item(n, "B", "AUTO_APPLY", excel=(500, "GM", None), proposals={"standard_size": None, "standard_uom": None, "standard_pack_size": "1"}) for n in range(100, 125)],
             "C": [item(n, "A", "REVIEW_REQUIRED", desc="PERSON PLEASE 500G") for n in range(200, 260)]}
    three = draw_more(items, 10, seed=5, samples=3)
    taken = [x["row_number"] for sampled in three for rows in sampled.values() for x in rows]
    assert len(taken) == len(set(taken)), "no row is in two samples"
    assert [{g: len(v) for g, v in sampled.items()} for sampled in three] == [{"A": 10, "B": 10, "C": 10}, {"A": 10, "B": 10, "C": 10}, {"A": 10, "B": 5, "C": 10}], \
        "a group too small for every sample gives the last one what is left"
    assert draw_more(items, 10, seed=5, samples=3) == three and draw_more(items, 10, seed=6, samples=3) != three
    # Samples added later never reuse a row of a sample already checked.
    first = {x["row_number"] for rows in three[0].values() for x in rows}
    later = draw_more(items, 10, seed=9, samples=2, used=first)
    assert not first & {x["row_number"] for sampled in later for rows in sampled.values() for x in rows}

    rows = KlmBlindTestService(MockKlmChecker(), 3).run_samples(list(zip([1, 2, 3], three)))
    assert sorted({r["sample"] for r in rows}) == [1, 2, 3]
    test = {"job_id": "j", "status": "READY", "size": 10, "seed": 5, "seeds": {"1": 5, "2": 5, "3": 5},
            "populations": {"A": 60, "B": 25, "C": 60}, "rows": rows, "checker_version": "v"}
    assert sample_numbers(test) == [1, 2, 3]
    assert summarize(test, 3)["overall"]["sample"] == 25 and summarize(test)["overall"]["sample"] == 85
    lines, _, header_row = lines_of(test, sample=2)
    assert lines[0][0].startswith("K L M blind test · sample 2 of 3 · 30 rows") and "share no row" in lines[1][0]
    assert len(lines) - header_row == 30 and sheet_name(2) == "KLM Blind Test 2"
    # Rows of a test stored before samples existed count as sample 1.
    old = {**test, "rows": [{k: v for k, v in r.items() if k != "sample"} for r in rows[:30]]}
    assert sample_numbers(old) == [1] and len(lines_of(old, sample=1)[0]) - lines_of(old, sample=1)[2] == 30
