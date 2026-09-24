from pathlib import Path

from openpyxl import Workbook

from app.agents.mock_provider import MockInferenceProvider
from app.agents.provider import (
    Evidence,
    InferenceResponse,
    InferenceResult,
    ObservedMeasurement,
    ProviderMetadata,
    InvalidInferenceResponseError,
)
from app.rules.registry import load_default_registry
from app.services.excel_reader import FIELD_MAP, INPUT_SHEET, PURGE_HEADERS
from app.services.processor import JobProcessor
from app.storage.local import LocalFileStorage


class FakeRepositories:
    def __init__(self, job):
        self.job = job
        self.items = []
        self.progress_updates = []

    def get_job(self, job_id):
        return self.job if self.job["job_id"] == job_id else None

    def update_job(self, job_id, values):
        for key, value in values.items():
            if "." not in key:
                self.job[key] = value
        if "progress" in values:
            self.progress_updates.append(dict(values["progress"]))

    def replace_items(self, documents):
        self.items.extend(documents)


def create_fixture(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = INPUT_SHEET
    headers = sorted(set(FIELD_MAP.values()) | PURGE_HEADERS)
    sheet.append(headers)

    def add(**values):
        sheet.append([values.get(header) for header in headers])

    add(Item_no="000001", Department="03_Grocery 2")  # purged
    add(Item_no="000002", Department="03_Grocery 2", item_desc_eng="JUICE 500ML", **{
        "Standardize Unit Size": 500, "Standardize UOM": "ML", "Standardize Pack Size": 1,
    })
    add(Item_no="000003", Department="03_Grocery 2", item_desc_eng="FLOUR", item_size_value=1, item_size_unit="KG", **{
        "Standardize Unit Size": 1, "Standardize UOM": "KG",
    })
    add(Item_no="000004", Department="06_Dairy & Frozen", item_desc_eng="GREEN TEA 1 KG")
    add(Item_no="000005", Department="Other", item_desc_eng="OUT OF SCOPE")
    workbook.save(path)


def test_processor_routes_rules_and_ai_without_crossing_paths(tmp_path: Path):
    storage = LocalFileStorage(tmp_path / "files", 10_000_000)
    job_id = "abc123"
    fixture = tmp_path / "fixture.xlsx"
    create_fixture(fixture)
    with fixture.open("rb") as source:
        storage.save_input(job_id, source)
    repository = FakeRepositories({
        "job_id": job_id,
        "selected_departments": ["03_Grocery 2", "06_Dairy & Frozen"],
        "snapshot_label": None,
    })
    processor = JobProcessor(repository, storage, load_default_registry(), MockInferenceProvider())
    processor.process(job_id)

    assert repository.job["status"] == "READY_FOR_REVIEW"
    assert repository.job["stats"] == {
        "workbook_rows": 5, "department_rows": 4, "purged": 1, "live": 3,
        "group_a": 1, "group_b": 1, "group_b1": 1, "group_b2": 0,
        "group_b3": 0, "group_c": 1, "validation_review": 0,
        "group_a_validation_warnings": 0,
        "data_shape_error": 0, "discrepancies": 0,
        "discrepancy_bilingual_measurement_conflicts": 0,
        "discrepancy_bilingual_count_conflicts": 0,
        "pack_existing_valid": 0, "pack_normalized_existing": 0,
        "pack_deterministic_proposed": 0, "pack_agent_proposed": 0,
        "pack_agent_declined": 1, "pack_conflict": 0, "pack_not_found": 1,
        "pack_agent_error": 0, "pack_agent_disabled": 0,
        "pack_invalid_existing": 0,
        "group_b_pack_deterministic_proposed": 0,
        "group_b_pack_agent_proposed": 0,
        "group_c_pack_deterministic_proposed": 0,
        "group_c_pack_agent_proposed": 0,
    }
    by_item = {item["item_no"]: item for item in repository.items}
    assert by_item["000003"]["method"] == "RULE"
    assert by_item["000003"]["field_proposals"]["standard_size"] == "1000"
    assert by_item["000004"]["method"] == "AI_INFERENCE"
    assert by_item["000004"]["reason_code"] == "NOT_IN_DESCRIPTION"
    assert by_item["000004"]["ai_provenance"]["provider"] == "mock"
    assert repository.job["quality"]["engine"]["processed_with_guards"] is True
    assert repository.job["quality"]["engine"]["guards_version"]


class ExtractingProvider:
    async def infer(self, request):
        return InferenceResponse(
            result=InferenceResult(
                status="PROPOSAL",
                measurement=ObservedMeasurement(
                    value="1",
                    uom="KG",
                    field="item_desc_eng",
                    fragment="1 KG",
                ),
                confidence="HIGH",
                reason_code="EXPLICIT_MEASUREMENT",
            ),
            metadata=ProviderMetadata(
                provider="google-adk",
                agent_name="test-agent",
                agent_version="1",
            ),
        )


def test_ai_observation_is_converted_only_by_deterministic_rule_engine(tmp_path: Path):
    storage = LocalFileStorage(tmp_path / "files", 10_000_000)
    job_id = "def456"
    fixture = tmp_path / "fixture.xlsx"
    create_fixture(fixture)
    with fixture.open("rb") as source:
        storage.save_input(job_id, source)
    repository = FakeRepositories({
        "job_id": job_id,
        "selected_departments": ["03_Grocery 2", "06_Dairy & Frozen"],
        "snapshot_label": None,
    })

    JobProcessor(repository, storage, load_default_registry(), ExtractingProvider()).process(job_id)

    item = next(item for item in repository.items if item["item_no"] == "000004")
    assert item["ai_observation"] == {
        "value": "1", "uom": "KG", "field": "item_desc_eng", "fragment": "1 KG", "role": None,
    }
    assert item["ai_raw"]["status"] == "PROPOSAL"
    assert item["confidence"] == "MEDIUM"  # derived: one clean source, not the model's "HIGH"
    assert repository.job["ai_usage"]["calls"] == 1
    assert item["field_proposals"]["standard_size"] == "1000"
    assert item["field_proposals"]["standard_uom"] == "GM"
    assert item["rule"]["rule_id"] == "KG_TO_GM"


def test_group_a_gate_normalizes_aliases_and_rejects_invalid_values(tmp_path: Path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = INPUT_SHEET
    headers = sorted(set(FIELD_MAP.values()) | PURGE_HEADERS)
    sheet.append(headers)

    def add(**values):
        sheet.append([values.get(header) for header in headers])

    add(Item_no="000101", Department="03_Grocery 2", item_desc_eng="JUICE", **{
        "Standardize Unit Size": 500.5,
        "Standardize UOM": "ml",
        "Standardize Pack Size": 1,
    })
    add(Item_no="000102", Department="03_Grocery 2", item_desc_eng="JUICE", **{
        "Standardize Unit Size": 0,
        "Standardize UOM": "ML",
        "Standardize Pack Size": 1,
    })
    add(Item_no="000103", Department="03_Grocery 2", item_desc_eng="JUICE", **{
        "Standardize Unit Size": "500.5",
        "Standardize UOM": "ML",
        "Standardize Pack Size": 1,
    })
    fixture = tmp_path / "group-a-gate.xlsx"
    workbook.save(fixture)

    storage = LocalFileStorage(tmp_path / "files", 10_000_000)
    job_id = "aabbccdd"
    with fixture.open("rb") as source:
        storage.save_input(job_id, source)
    repository = FakeRepositories({
        "job_id": job_id,
        "selected_departments": ["03_Grocery 2", "06_Dairy & Frozen"],
        "snapshot_label": None,
    })

    JobProcessor(repository, storage, load_default_registry(), MockInferenceProvider()).process(job_id)

    by_item = {item["item_no"]: item for item in repository.items}
    normalized = by_item["000101"]
    assert normalized["group"] == "B"
    assert normalized["reason_code"] == "STANDARD_FIELDS_NORMALIZATION"
    assert normalized["field_proposals"] == {
        "standard_size": None,
        "standard_uom": "ML",
        "standard_pack_size": None,
    }
    assert normalized["rule"]["rounding_decimals"] is None
    assert normalized["original"]["standard_uom"] == "ml"
    assert normalized["validation"]["status"] == "AUTO_FIX"

    numeric_text = by_item["000103"]
    assert numeric_text["group"] == "B"
    assert numeric_text["field_proposals"] == {
        "standard_size": "500.5",
        "standard_uom": None,
        "standard_pack_size": None,
    }
    assert numeric_text["rule"]["rounding_decimals"] is None

    invalid = by_item["000102"]
    assert invalid["group"] == "DATA_SHAPE_ERROR"
    assert invalid["reason_code"] == "GROUP_A_VALIDATION_INVALID"
    assert invalid["validation"]["status"] == "INVALID"


class PackAwareProvider:
    def __init__(self):
        self.requests = []

    async def infer(self, request):
        self.requests.append(request)
        if request.task == "PACK_ONLY":
            return InferenceResponse(
                result=InferenceResult(
                    status="PACK_PROPOSAL",
                    pack_size=4,
                    pack_evidence=Evidence(field="item_desc_eng", fragment="4'S"),
                    confidence="HIGH",
                    reason_code="EXPLICIT_PACK_COUNT",
                ),
                metadata=ProviderMetadata(
                    provider="mock", agent_name="pack-test", agent_version="1"
                ),
            )
        return InferenceResponse(
            result=InferenceResult(
                status="PROPOSAL",
                measurement=ObservedMeasurement(
                    value=500,
                    uom="ML",
                    field="item_desc_eng",
                    fragment="500ML",
                ),
                confidence="HIGH",
                reason_code="EXPLICIT_MEASUREMENT",
            ),
            metadata=ProviderMetadata(
                provider="mock", agent_name="measurement-test", agent_version="1"
            ),
        )


def test_b_and_c_share_evidence_first_pack_pipeline(tmp_path: Path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = INPUT_SHEET
    headers = sorted(set(FIELD_MAP.values()) | PURGE_HEADERS)
    sheet.append(headers)

    def add(**values):
        sheet.append([values.get(header) for header in headers])

    add(Item_no="000201", Department="03_Grocery 2", item_desc_eng="JUICE 6 X 500ML",
        item_size_value=500, item_size_unit="ML")
    add(Item_no="000202", Department="03_Grocery 2", item_desc_eng="JUICE ASSORTED 4'S 500ML",
        item_size_value=500, item_size_unit="ML")
    add(Item_no="000203", Department="06_Dairy & Frozen", item_desc_eng="WATER 6 X 500ML")
    fixture = tmp_path / "pack-pipeline.xlsx"
    workbook.save(fixture)

    storage = LocalFileStorage(tmp_path / "files", 10_000_000)
    job_id = "facefeed"
    with fixture.open("rb") as source:
        storage.save_input(job_id, source)
    repository = FakeRepositories({
        "job_id": job_id,
        "selected_departments": ["03_Grocery 2", "06_Dairy & Frozen"],
        "snapshot_label": None,
    })
    provider = PackAwareProvider()

    JobProcessor(repository, storage, load_default_registry(), provider).process(job_id)

    by_item = {item["item_no"]: item for item in repository.items}
    deterministic_b = by_item["000201"]
    assert deterministic_b["field_proposals"]["standard_pack_size"] == "6"
    assert deterministic_b["pack_result"]["status"] == "DETERMINISTIC_PROPOSAL"
    assert deterministic_b["field_provenance"]["standard_pack_size"]["method"] == "RULE"

    agent_b = by_item["000202"]
    assert agent_b["field_proposals"]["standard_pack_size"] == "4"
    assert agent_b["pack_result"]["status"] == "AGENT_PROPOSAL"
    assert agent_b["field_provenance"]["standard_pack_size"]["method"] == "AI_INFERENCE"
    assert agent_b["field_provenance"]["standard_size"]["method"] == "RULE"

    deterministic_c = by_item["000203"]
    assert deterministic_c["field_proposals"] == {
        "standard_size": "500",
        "standard_uom": "ML",
        "standard_pack_size": "6",
    }
    assert deterministic_c["pack_result"]["status"] == "DETERMINISTIC_PROPOSAL"

    assert [request.task for request in provider.requests].count("PACK_ONLY") == 1
    assert [request.task for request in provider.requests].count("MEASUREMENT_AND_PACK") == 1


def _run_single_b_pack_case(tmp_path: Path, provider, *, enabled: bool):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = INPUT_SHEET
    headers = sorted(set(FIELD_MAP.values()) | PURGE_HEADERS)
    sheet.append(headers)
    values = {
        "Item_no": "000301",
        "Department": "03_Grocery 2",
        "item_desc_eng": "JUICE ASSORTED 4'S 500ML",
        "item_size_value": 500,
        "item_size_unit": "ML",
    }
    sheet.append([values.get(header) for header in headers])
    fixture = tmp_path / "single-pack.xlsx"
    workbook.save(fixture)
    storage = LocalFileStorage(tmp_path / "files", 10_000_000)
    job_id = "deadbeef"
    with fixture.open("rb") as source:
        storage.save_input(job_id, source)
    repository = FakeRepositories({
        "job_id": job_id,
        "selected_departments": ["03_Grocery 2", "06_Dairy & Frozen"],
        "snapshot_label": None,
    })
    JobProcessor(
        repository, storage, load_default_registry(), provider,
        pack_size_inference_enabled=enabled,
    ).process(job_id)
    return repository.items[0]


class MustNotCallProvider:
    async def infer(self, request):
        raise AssertionError("provider must not be called")


def test_pack_agent_feature_flag_does_not_affect_b_conversion(tmp_path: Path):
    item = _run_single_b_pack_case(tmp_path, MustNotCallProvider(), enabled=False)
    assert item["field_proposals"]["standard_size"] == "500"
    assert item["field_proposals"]["standard_uom"] == "ML"
    assert item["field_proposals"]["standard_pack_size"] is None
    assert item["pack_result"]["status"] == "AGENT_DISABLED"


class FailingPackProvider:
    async def infer(self, request):
        raise RuntimeError("provider unavailable")


def test_pack_agent_failure_isolated_from_b_conversion(tmp_path: Path):
    item = _run_single_b_pack_case(tmp_path, FailingPackProvider(), enabled=True)
    assert item["field_proposals"]["standard_size"] == "500"
    assert item["field_proposals"]["standard_uom"] == "ML"
    assert item["field_proposals"]["standard_pack_size"] is None
    assert item["pack_result"]["status"] == "AGENT_ERROR"
    assert item["pack_result"]["reason_code"] == "PACK_AGENT_ERROR"


class InvalidPackProvider:
    async def infer(self, request):
        raise InvalidInferenceResponseError(
            "repair exhausted",
            validation_errors=["first invalid answer", "second invalid answer"],
        )


def test_pack_agent_records_both_invalid_response_attempts(tmp_path: Path):
    item = _run_single_b_pack_case(tmp_path, InvalidPackProvider(), enabled=True)

    assert item["pack_result"]["status"] == "AGENT_ERROR"
    assert item["pack_result"]["validation_attempts"] == [
        {"attempt": 1, "error": "first invalid answer"},
        {"attempt": 2, "error": "second invalid answer"},
    ]


def test_progress_is_reported_while_the_job_runs(tmp_path: Path):
    storage = LocalFileStorage(tmp_path / "files", 10_000_000)
    job_id = "abc123"
    fixture = tmp_path / "fixture.xlsx"
    create_fixture(fixture)
    with fixture.open("rb") as source:
        storage.save_input(job_id, source)
    repository = FakeRepositories({
        "job_id": job_id,
        "selected_departments": ["03_Grocery 2", "06_Dairy & Frozen"],
        "snapshot_label": None,
    })
    JobProcessor(repository, storage, load_default_registry(), MockInferenceProvider()).process(job_id)

    updates = repository.progress_updates
    stages = list(dict.fromkeys(update["stage"] for update in updates))
    assert stages == [
        "PROFILING", "PROCESSING_RULES", "PROCESSING_DESCRIPTIONS",
        "CHECKING_DISCREPANCIES", "SAVING_RESULTS", "READY_FOR_REVIEW",
    ]
    # Every stage with a known size reports its final row before the next begins.
    rules = [update for update in updates if update["stage"] == "PROCESSING_RULES"]
    assert rules[0]["processed"] == 0 and rules[-1] == {
        "stage": "PROCESSING_RULES", "processed": 4, "total": 4, "unit": "ROWS", "percent": 55,
    }
    agent = [update for update in updates if update["stage"] == "PROCESSING_DESCRIPTIONS"]
    assert agent[-1]["processed"] == agent[-1]["total"] == 1
    assert agent[-1]["unit"] == "AGENT_CALLS"
    percents = [update["percent"] for update in updates]
    assert percents == sorted(percents) and percents[-1] == 100


def test_count_conflict_requires_review_without_changing_the_group(tmp_path: Path):
    storage = LocalFileStorage(tmp_path / "files", 10_000_000)
    job_id = "abc123"
    fixture = tmp_path / "fixture.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = INPUT_SHEET
    headers = sorted(set(FIELD_MAP.values()) | PURGE_HEADERS)
    sheet.append(headers)
    values = {
        "Item_no": "226167", "Department": "06_Dairy & Frozen",
        "item_desc_eng": "MINI DAIFUKU STRAWBERRY 9S", "item_desc_local_lang": "迷你大福草莓雪米糍3件裝",
        "Standardize Unit Size": 30, "Standardize UOM": "ML", "Standardize Pack Size": 9,
    }
    sheet.append([values.get(header) for header in headers])
    workbook.save(fixture)
    with fixture.open("rb") as source:
        storage.save_input(job_id, source)
    repository = FakeRepositories({
        "job_id": job_id, "selected_departments": ["06_Dairy & Frozen"], "snapshot_label": None,
    })
    JobProcessor(repository, storage, load_default_registry(), MockInferenceProvider()).process(job_id)

    item = repository.items[0]
    assert item["group"] == "A"
    assert item["discrepancy"]["flagged"] is True
    assert item["application_policy"] == "REVIEW_REQUIRED"
    assert item["review"]["overall_status"] == "PENDING"
    assert [finding["code"] for finding in item["findings"]] == ["PACK_COUNT_CONFLICT"]
    # A conflict never proposes or applies a value.
    assert all(change["proposed"] is None for change in item["changes"])
    assert repository.job["stats"]["discrepancy_bilingual_count_conflicts"] == 1


def run_rows(tmp_path, rows, provider=None, departments=("03_Grocery 2", "06_Dairy & Frozen")):
    storage = LocalFileStorage(tmp_path / "files", 10_000_000)
    fixture = tmp_path / "fixture.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = INPUT_SHEET
    headers = sorted(set(FIELD_MAP.values()) | PURGE_HEADERS)
    sheet.append(headers)
    for values in rows:
        sheet.append([values.get(header) for header in headers])
    workbook.save(fixture)
    with fixture.open("rb") as source:
        storage.save_input("abc123", source)
    repository = FakeRepositories({
        "job_id": "abc123", "selected_departments": list(departments), "snapshot_label": None,
    })
    JobProcessor(repository, storage, load_default_registry(), provider or MockInferenceProvider()).process("abc123")
    return repository


def validated(number, category, size, uom):
    return {"Item_no": f"9{number:05d}", "Department": "06_Dairy & Frozen", "Category": category,
            "item_desc_eng": "FILLER", "Standardize Unit Size": size, "Standardize UOM": uom,
            "Standardize Pack Size": 1}


def test_ounce_is_a_fluid_ounce_in_a_liquid_category_and_a_question_in_a_mixed_one(tmp_path: Path):
    rows = [validated(n, "Fresh Milk", 1000, "ML") for n in range(20)]
    rows += [validated(100 + n, "Sauces", 300, "ML" if n % 2 else "GM") for n in range(20)]
    rows += [validated(200 + n, "Flour", 1000, "GM") for n in range(20)]
    for number, category in (("000001", "Fresh Milk"), ("000002", "Sauces"), ("000003", "Flour")):
        rows.append({"Item_no": number, "Department": "06_Dairy & Frozen", "Category": category,
                     "item_desc_eng": "PRODUCT", "item_size_value": 48, "item_size_unit": "OZ",
                     "Standardize Unit Size": 48, "Standardize UOM": "OZ", "Standardize Pack Size": 1})
    items = {item["item_no"]: item for item in run_rows(tmp_path, rows).items}

    milk, sauce, flour = items["000001"], items["000002"], items["000003"]
    assert (milk["field_proposals"]["standard_size"], milk["field_proposals"]["standard_uom"]) == ("1420", "ML")
    assert milk["rule"]["rule_id"] == "FLOZ_TO_ML" and milk["application_policy"] == "AUTO_APPLY"
    assert [f["code"] for f in milk["findings"]] == ["OUNCE_READ_AS_FLUID"]

    # Mixed category: keep the weight as the suggestion, but a person chooses.
    assert (sauce["field_proposals"]["standard_size"], sauce["field_proposals"]["standard_uom"]) == ("1361", "GM")
    assert sauce["application_policy"] == "REVIEW_REQUIRED"
    assert "1361 GM" in sauce["findings"][0]["human_reason"] and "1420 ML" in sauce["findings"][0]["human_reason"]

    assert (flour["field_proposals"]["standard_uom"], flour["application_policy"], flour["findings"]) == ("GM", "AUTO_APPLY", [])


class RoleProvider:
    """Answers like agent contract v3: every reading carries a role."""

    def __init__(self, role, pack_role=None):
        self.role, self.pack_role, self.requests = role, pack_role, []

    async def infer(self, request):
        self.requests.append(request)
        result = InferenceResult(
            status="PROPOSAL", rationale="test",
            measurement=ObservedMeasurement(value="1", uom="L", field="item_desc_eng", fragment="1L", role=self.role),
            pack_size="8" if self.pack_role else None,
            pack_evidence=Evidence(field="item_desc_eng", fragment="8 PACK") if self.pack_role else None,
            pack_role=self.pack_role,
        )
        return InferenceResponse(result=result, metadata=ProviderMetadata(
            provider="google-adk", agent_name="a", agent_version="2.0.0", input_tokens=900, output_tokens=60,
        ))


def blank_row(description, category="Baking Aids"):
    return [{"Item_no": "000777", "Department": "03_Grocery 2", "Category": category, "item_desc_eng": description}]


def test_a_capacity_reading_never_becomes_a_size_whatever_the_model_recommends(tmp_path: Path):
    provider = RoleProvider("CAPACITY_OR_RANGE")
    repository = run_rows(tmp_path, blank_row("1L MICROWAVE BOX (8 PACK)"), provider)
    item = repository.items[0]

    assert item["field_proposals"] == {"standard_size": None, "standard_uom": None, "standard_pack_size": None}
    assert item["reason_code"] == "AI_MEASUREMENT_NOT_PRODUCT_SIZE"
    assert item["application_policy"] == "UNRESOLVED"
    reason = next(f["human_reason"] for f in item["findings"] if f["code"] == "AI_MEASUREMENT_NOT_PRODUCT_SIZE")
    assert "“1L”" in reason and "capacity" in reason
    # D3: the agent is given the category as context, and nothing it must not see.
    sent = provider.requests[0].model_dump(exclude_none=True)
    assert sent["category"] == "Baking Aids"
    assert not {"standard_size", "legacy_size", "product_description"} & set(sent)
    assert repository.job["ai_usage"] == {"calls": 1, "input_tokens": 900, "output_tokens": 60}


def test_a_contents_count_is_not_written_as_the_pack_size(tmp_path: Path):
    contents = run_rows(tmp_path / "a", blank_row("JUICE 1L (8 PACK)"), RoleProvider("NET_CONTENT_TOTAL", "CONTENTS")).items[0]
    sellable = run_rows(tmp_path / "b", blank_row("JUICE 1L (8 PACK)"), RoleProvider("NET_CONTENT_UNIT", "SELLABLE_PACK")).items[0]

    assert contents["field_proposals"]["standard_size"] == "1000"
    assert contents["field_proposals"]["standard_pack_size"] is None
    assert sellable["field_proposals"]["standard_pack_size"] == "8"
    # An AI-read pack size is a suggestion for a person, never an automatic write.
    assert sellable["application_policy"] == "REVIEW_REQUIRED"
    assert "AI_PACK_NEEDS_CONFIRMATION" in [f["code"] for f in sellable["findings"]]
    pack = next(c for c in sellable["changes"] if c["field"] == "standard_pack_size")
    assert (pack["proposed"], pack["final"]) == ("8", None)


class OnePieceVoucherProvider:
    """Reads the pack count 1 from “1PC”, the way the live reader does on vouchers."""
    async def infer(self, request):
        from app.agents.provider import Evidence, InferenceResponse, InferenceResult, ProviderMetadata
        if request.task == "PACK_ONLY":
            result = InferenceResult(status="PACK_PROPOSAL", pack_size=1, pack_role="SELLABLE_PACK",
                                     pack_evidence=Evidence(field="item_desc_local_lang", fragment="1PC"), rationale="one voucher", confidence="HIGH")
        else:
            result = InferenceResult(status="NOT_IN_DESCRIPTION", rationale="no size")
        return InferenceResponse(result=result, metadata=ProviderMetadata(provider="mock", agent_name="t", agent_version="1", prompt_version="v", prompt_sha256="x", model_id="m", adk_version="1", attempt_count=1, latency_ms=1))


def test_a_pack_count_of_one_that_matches_the_legacy_needs_no_confirmation(tmp_path: Path):
    repository = run_rows(tmp_path, [{
        "Item_no": "231076", "Department": "03_Grocery 2", "item_desc_eng": "WATER CHESTNUT P V",
        "item_desc_local_lang": "利苑清香馬蹄糕禮券1PC", "item_size_value": 1, "item_size_unit": "PC",
    }], provider=OnePieceVoucherProvider())
    item = repository.items[0]
    assert item["group"] == "B"
    assert item["field_proposals"]["standard_size"] == "1" and item["field_proposals"]["standard_uom"] == "EA"
    assert item["field_proposals"]["standard_pack_size"] == "1"
    assert "AI_PACK_NEEDS_CONFIRMATION" not in {g["code"] for g in item["guards"]}
    assert item["application_policy"] == "AUTO_APPLY"
